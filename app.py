from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, Response, jsonify
)
import csv
import io
import os
import re
import sqlite3
import psycopg
from psycopg.rows import dict_row
from psycopg.errors import IntegrityError as PostgresIntegrityError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)
app.secret_key = os.environ.get(
    "SECRET_KEY",
    "change-this-to-a-long-private-secret-key-before-deployment"
)

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
SQLITE_DATABASE = os.path.join(BASE_DIR, "office_day.db")
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
USING_POSTGRES = bool(DATABASE_URL)

PURPOSES = [
    "Submit Report",
    "Regular Office Day",
    "Office Day Requested by Sales Department",
    "Other"
]

COMPANIES = [
    "CONCEPT CLOTHING CO., INC.",
    "CORPORATE APPAREL, INC.",
    "CEO GROUP OF COMPANIES"
]


class DBConnection:
    """Small compatibility wrapper so the existing Flask routes work on SQLite or PostgreSQL."""
    def __init__(self):
        if USING_POSTGRES:
            self.raw = psycopg.connect(DATABASE_URL, row_factory=dict_row)
        else:
            self.raw = sqlite3.connect(SQLITE_DATABASE)
            self.raw.row_factory = sqlite3.Row
            self.raw.execute("PRAGMA foreign_keys = ON")

    def execute(self, sql, params=()):
        if USING_POSTGRES:
            # Existing app uses SQLite-style ? placeholders.
            sql = sql.replace("?", "%s")
            # SQLite-specific case-insensitive collation is replaced by PostgreSQL comparisons.
            sql = re.sub(r"\s+COLLATE\s+NOCASE", "", sql, flags=re.I)
        return self.raw.execute(sql, params)

    def commit(self):
        return self.raw.commit()

    def rollback(self):
        return self.raw.rollback()

    def close(self):
        return self.raw.close()


def get_db_connection():
    return DBConnection()


def table_columns(connection, table_name):
    if USING_POSTGRES:
        rows = connection.execute("""
            SELECT column_name AS name
            FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = ?
        """, (table_name,)).fetchall()
        return {row["name"] for row in rows}

    return {
        row["name"]
        for row in connection.execute(f"PRAGMA table_info({table_name})").fetchall()
    }


def table_exists(connection, table_name):
    if USING_POSTGRES:
        row = connection.execute("""
            SELECT 1
            FROM information_schema.tables
            WHERE table_schema = 'public' AND table_name = ?
        """, (table_name,)).fetchone()
    else:
        row = connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
            (table_name,)
        ).fetchone()
    return row is not None


def migrate_users(connection):
    id_type = "SERIAL PRIMARY KEY" if USING_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS users (
            id {id_type},
            full_name TEXT NOT NULL,
            username TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('ADMIN', 'HR', 'SALES')),
            status TEXT NOT NULL DEFAULT 'Active'
                CHECK(status IN ('Active', 'Inactive')),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    columns = table_columns(connection, "users")
    if "status" not in columns:
        connection.execute(
            "ALTER TABLE users ADD COLUMN status TEXT NOT NULL DEFAULT 'Active'"
        )



def migrate_user_brands(connection):
    connection.execute("""
        CREATE TABLE IF NOT EXISTS user_brands (
            user_id INTEGER NOT NULL,
            brand_name TEXT NOT NULL,
            PRIMARY KEY (user_id, brand_name),
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        )
    """)

def migrate_stores(connection):
    id_type = "SERIAL PRIMARY KEY" if USING_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS stores (
            id {id_type},
            store_name TEXT NOT NULL UNIQUE,
            status TEXT NOT NULL DEFAULT 'Active'
                CHECK(status IN ('Active', 'Inactive')),
            remarks TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    columns = table_columns(connection, "stores")
    if "status" not in columns:
        connection.execute(
            "ALTER TABLE stores ADD COLUMN status TEXT NOT NULL DEFAULT 'Active'"
        )
    if "remarks" not in columns:
        connection.execute("ALTER TABLE stores ADD COLUMN remarks TEXT")
    if "created_at" not in columns:
        connection.execute("ALTER TABLE stores ADD COLUMN created_at TIMESTAMP")


def migrate_brands(connection):
    id_type = "SERIAL PRIMARY KEY" if USING_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS brands (
            id {id_type},
            brand_name TEXT NOT NULL UNIQUE,
            status TEXT NOT NULL DEFAULT 'Active'
                CHECK(status IN ('Active', 'Inactive')),
            remarks TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)



def migrate_companies(connection):
    id_type = "SERIAL PRIMARY KEY" if USING_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS companies (
            id {id_type},
            company_name TEXT NOT NULL UNIQUE,
            status TEXT NOT NULL DEFAULT 'Active'
                CHECK(status IN ('Active', 'Inactive')),
            remarks TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    columns = table_columns(connection, "companies")
    if "status" not in columns:
        connection.execute(
            "ALTER TABLE companies ADD COLUMN status TEXT NOT NULL DEFAULT 'Active'"
        )
    if "remarks" not in columns:
        connection.execute("ALTER TABLE companies ADD COLUMN remarks TEXT")
    if "created_at" not in columns:
        connection.execute("ALTER TABLE companies ADD COLUMN created_at TIMESTAMP")


def migrate_employees(connection):
    id_type = "SERIAL PRIMARY KEY" if USING_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS employees (
            id {id_type},
            company_name TEXT NOT NULL,
            last_name TEXT NOT NULL,
            first_name TEXT NOT NULL,
            middle_name TEXT,
            position TEXT NOT NULL,
            brand TEXT NOT NULL,
            store_assignment TEXT NOT NULL,
            date_hired TEXT,
            status TEXT NOT NULL DEFAULT 'Active'
                CHECK(status IN ('Active', 'Inactive')),
            remarks TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)


def migrate_requests(connection):
    id_type = "SERIAL PRIMARY KEY" if USING_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS office_day_requests (
            id {id_type},
            employee_id INTEGER,
            employee_name TEXT NOT NULL,
            company_name TEXT,
            position TEXT,
            store_assignment TEXT NOT NULL,
            brand TEXT NOT NULL,
            requested_date TEXT NOT NULL,
            purpose TEXT NOT NULL,
            other_purpose TEXT,
            remarks TEXT,
            requested_by INTEGER NOT NULL,
            requested_by_name TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Pending HR Approval',
            approved_date TEXT,
            hr_remarks TEXT,
            hr_action_by TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (employee_id) REFERENCES employees(id),
            FOREIGN KEY (requested_by) REFERENCES users(id)
        )
    """)

    columns = table_columns(connection, "office_day_requests")
    additions = {
        "employee_id": "INTEGER",
        "company_name": "TEXT",
        "position": "TEXT",
        "hr_remarks": "TEXT"
    }
    for column, datatype in additions.items():
        if column not in columns:
            connection.execute(
                f"ALTER TABLE office_day_requests ADD COLUMN {column} {datatype}"
            )


def initialize_database():
    connection = get_db_connection()
    try:
        migrate_users(connection)
        migrate_user_brands(connection)
        migrate_stores(connection)
        migrate_brands(connection)
        migrate_companies(connection)
        migrate_employees(connection)
        migrate_requests(connection)

        default_brands = [
            "WALL STREET", "SAHARA", "CRITERION", "ULTIMO",
            "ARROW", "VAN HEUSEN", "IZOD"
        ]
        for brand in default_brands:
            if USING_POSTGRES:
                connection.execute("""
                    INSERT INTO brands (brand_name, status)
                    VALUES (?, 'Active')
                    ON CONFLICT (brand_name) DO NOTHING
                """, (brand,))
            else:
                connection.execute("""
                    INSERT OR IGNORE INTO brands (brand_name, status)
                    VALUES (?, 'Active')
                """, (brand,))

        user_count = connection.execute(
            "SELECT COUNT(*) AS total FROM users"
        ).fetchone()["total"]

        if user_count == 0:
            seed_users = [
                ("System Administrator", "admin", "admin12345", "ADMIN"),
                ("HR Administrator", "hradmin", "hr12345", "HR"),
                ("Sales Department", "salesuser", "sales12345", "SALES")
            ]
            for full_name, username, password, role in seed_users:
                connection.execute("""
                    INSERT INTO users (
                        full_name, username, password, role, status
                    )
                    VALUES (?, ?, ?, ?, 'Active')
                """, (
                    full_name, username,
                    generate_password_hash(password), role
                ))

        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def login_required(view_function):
    @wraps(view_function)
    def decorated_function(*args, **kwargs):
        if "user_id" not in session:
            flash("Please log in first.", "error")
            return redirect(url_for("login"))
        return view_function(*args, **kwargs)
    return decorated_function


def roles_required(*allowed_roles):
    def decorator(view_function):
        @wraps(view_function)
        def decorated_function(*args, **kwargs):
            if "user_id" not in session:
                flash("Please log in first.", "error")
                return redirect(url_for("login"))
            if session.get("role") not in allowed_roles:
                flash("You are not authorized to access that page.", "error")
                return redirect(url_for("dashboard"))
            return view_function(*args, **kwargs)
        return decorated_function
    return decorator



def get_user_brands(connection, user_id):
    rows = connection.execute("""
        SELECT brand_name
        FROM user_brands
        WHERE user_id=?
        ORDER BY brand_name
    """, (user_id,)).fetchall()
    return [row["brand_name"] for row in rows]


def replace_user_brands(connection, user_id, brand_names):
    connection.execute("DELETE FROM user_brands WHERE user_id=?", (user_id,))
    for brand_name in brand_names:
        connection.execute("""
            INSERT INTO user_brands (user_id, brand_name)
            VALUES (?, ?)
        """, (user_id, brand_name))


def employee_display_name(row):
    middle = (row["middle_name"] or "").strip()
    middle_initial = f" {middle[0].upper()}." if middle else ""
    return f'{row["last_name"].upper()}, {row["first_name"]}{middle_initial}'


def fetch_dashboard_counts(connection):
    return {
        "employees": connection.execute(
            "SELECT COUNT(*) AS count FROM employees WHERE status='Active'"
        ).fetchone()["count"],
        "pending": connection.execute(
            "SELECT COUNT(*) AS count FROM office_day_requests "
            "WHERE status='Pending HR Approval'"
        ).fetchone()["count"],
        "approved": connection.execute(
            "SELECT COUNT(*) AS count FROM office_day_requests "
            "WHERE status IN ('Approved', 'Adjusted Schedule')"
        ).fetchone()["count"],
        "users": connection.execute(
            "SELECT COUNT(*) AS count FROM users WHERE status='Active'"
        ).fetchone()["count"]
    }


@app.context_processor
def inject_helpers():
    return {
        "employee_display_name": employee_display_name,
        "PURPOSES": PURPOSES
    }


@app.route("/")
def home():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return redirect(url_for("dashboard"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        connection = get_db_connection()
        user = connection.execute("""
            SELECT * FROM users
            WHERE LOWER(username)=LOWER(?) AND status = 'Active'
        """, (username,)).fetchone()
        connection.close()

        if user and check_password_hash(user["password"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["full_name"] = user["full_name"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            return redirect(url_for("dashboard"))

        flash("Invalid username/password or inactive account.", "error")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    session.clear()
    flash("You have been logged out.", "success")
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard():
    role = session["role"]
    if role == "ADMIN":
        return redirect(url_for("admin_dashboard"))
    if role == "HR":
        return redirect(url_for("hr_dashboard"))
    return redirect(url_for("sales_dashboard"))


@app.route("/admin")
@roles_required("ADMIN")
def admin_dashboard():
    connection = get_db_connection()
    counts = fetch_dashboard_counts(connection)
    recent_requests = connection.execute("""
        SELECT * FROM office_day_requests ORDER BY created_at DESC LIMIT 8
    """).fetchall()
    connection.close()
    return render_template(
        "admin_dashboard.html",
        counts=counts,
        recent_requests=recent_requests
    )


@app.route("/hr")
@roles_required("ADMIN", "HR")
def hr_dashboard():
    connection = get_db_connection()
    counts = fetch_dashboard_counts(connection)
    office_requests = connection.execute("""
        SELECT * FROM office_day_requests
        ORDER BY
            CASE status
                WHEN 'Pending HR Approval' THEN 1
                WHEN 'Adjusted Schedule' THEN 2
                WHEN 'Approved' THEN 3
                ELSE 4
            END,
            created_at DESC
    """).fetchall()
    connection.close()
    return render_template(
        "hr_dashboard.html",
        counts=counts,
        office_requests=office_requests
    )


@app.route("/sales")
@roles_required("SALES")
def sales_dashboard():
    connection = get_db_connection()
    assigned_brands = get_user_brands(connection, session["user_id"])

    if assigned_brands:
        placeholders = ", ".join(["?"] * len(assigned_brands))
        employees = connection.execute(f"""
            SELECT * FROM employees
            WHERE status='Active'
              AND LOWER(brand) IN ({placeholders})
            ORDER BY last_name, first_name, middle_name
        """, [brand.lower() for brand in assigned_brands]).fetchall()
    else:
        employees = []

    office_requests = connection.execute("""
        SELECT * FROM office_day_requests
        WHERE requested_by = ?
        ORDER BY created_at DESC
    """, (session["user_id"],)).fetchall()
    connection.close()

    employee_data = {
        str(row["id"]): {
            "company_name": row["company_name"],
            "position": row["position"],
            "brand": row["brand"],
            "store_assignment": row["store_assignment"]
        }
        for row in employees
    }
    return render_template(
        "sales_dashboard.html",
        employees=employees,
        office_requests=office_requests,
        employee_data=employee_data,
        purposes=PURPOSES,
        assigned_brands=assigned_brands
    )

@app.route("/sales/submit-request", methods=["POST"])
@roles_required("SALES")
def submit_request():
    employee_id = request.form.get("employee_id", "").strip()
    requested_date = request.form.get("requested_date", "").strip()
    purpose = request.form.get("purpose", "").strip()
    other_purpose = request.form.get("other_purpose", "").strip()
    remarks = request.form.get("remarks", "").strip()

    if not employee_id or not requested_date or purpose not in PURPOSES:
        flash("Please complete all required fields.", "error")
        return redirect(url_for("sales_dashboard"))

    if purpose == "Other" and not other_purpose:
        flash("Please specify the other purpose.", "error")
        return redirect(url_for("sales_dashboard"))

    connection = get_db_connection()
    assigned_brands = get_user_brands(connection, session["user_id"])
    employee = connection.execute("""
        SELECT * FROM employees WHERE id=? AND status='Active'
    """, (employee_id,)).fetchone()

    if employee is None:
        connection.close()
        flash("Selected employee is unavailable or inactive.", "error")
        return redirect(url_for("sales_dashboard"))

    if not assigned_brands or employee["brand"].strip().lower() not in {
        brand.strip().lower() for brand in assigned_brands
    }:
        connection.close()
        flash("You are not authorized to submit requests for this employee's brand.", "error")
        return redirect(url_for("sales_dashboard"))

    display_name = employee_display_name(employee)
    connection.execute("""
        INSERT INTO office_day_requests (
            employee_id, employee_name, company_name, position,
            store_assignment, brand, requested_date, purpose,
            other_purpose, remarks, requested_by, requested_by_name,
            status
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pending HR Approval')
    """, (
        employee["id"], display_name, employee["company_name"],
        employee["position"], employee["store_assignment"],
        employee["brand"], requested_date, purpose, other_purpose,
        remarks, session["user_id"], session["full_name"]
    ))
    connection.commit()
    connection.close()

    flash("Office Day request submitted successfully.", "success")
    return redirect(url_for("sales_dashboard"))


@app.route("/hr/request/<int:request_id>/approve", methods=["POST"])
@roles_required("ADMIN", "HR")
def approve_request(request_id):
    hr_remarks = request.form.get("hr_remarks", "").strip()
    connection = get_db_connection()
    cursor = connection.execute("""
        UPDATE office_day_requests
        SET status='Approved',
            approved_date=requested_date,
            hr_remarks=?,
            hr_action_by=?
        WHERE id=?
    """, (hr_remarks, session["full_name"], request_id))
    connection.commit()
    connection.close()

    flash(
        "Office Day request approved." if cursor.rowcount
        else "Office Day request not found.",
        "success" if cursor.rowcount else "error"
    )
    return redirect(url_for("hr_dashboard"))


@app.route("/hr/request/<int:request_id>/adjust", methods=["POST"])
@roles_required("ADMIN", "HR")
def adjust_request(request_id):
    adjusted_date = request.form.get("adjusted_date", "").strip()
    hr_remarks = request.form.get("hr_remarks", "").strip()

    if not adjusted_date or not hr_remarks:
        flash("Adjusted date and HR remarks are required.", "error")
        return redirect(url_for("hr_dashboard"))

    connection = get_db_connection()
    cursor = connection.execute("""
        UPDATE office_day_requests
        SET status='Adjusted Schedule',
            approved_date=?,
            hr_remarks=?,
            hr_action_by=?
        WHERE id=?
    """, (adjusted_date, hr_remarks, session["full_name"], request_id))
    connection.commit()
    connection.close()

    flash(
        "Adjusted Office Day schedule saved." if cursor.rowcount
        else "Office Day request not found.",
        "success" if cursor.rowcount else "error"
    )
    return redirect(url_for("hr_dashboard"))


@app.route("/hr/request/<int:request_id>/reject", methods=["POST"])
@roles_required("ADMIN", "HR")
def reject_request(request_id):
    hr_remarks = request.form.get("hr_remarks", "").strip()
    if not hr_remarks:
        flash("HR remarks are required when rejecting a request.", "error")
        return redirect(url_for("hr_dashboard"))

    connection = get_db_connection()
    cursor = connection.execute("""
        UPDATE office_day_requests
        SET status='Rejected',
            approved_date=NULL,
            hr_remarks=?,
            hr_action_by=?
        WHERE id=?
    """, (hr_remarks, session["full_name"], request_id))
    connection.commit()
    connection.close()

    flash(
        "Office Day request rejected." if cursor.rowcount
        else "Office Day request not found.",
        "success" if cursor.rowcount else "error"
    )
    return redirect(url_for("hr_dashboard"))


@app.route("/requests/<int:request_id>/delete", methods=["POST"])
@login_required
def delete_request(request_id):
    return_to = request.form.get("return_to", "dashboard").strip()
    allowed_destinations = {
        "hr_dashboard": "hr_dashboard",
        "sales_dashboard": "sales_dashboard",
        "reports": "reports"
    }
    destination = allowed_destinations.get(return_to, "dashboard")

    connection = get_db_connection()
    try:
        office_request = connection.execute(
            """
            SELECT id, requested_by, status
            FROM office_day_requests
            WHERE id = ?
            """,
            (request_id,)
        ).fetchone()

        if office_request is None:
            flash("Office Day request not found.", "error")
            return redirect(url_for(destination))

        if session.get("role") == "SALES":
            if office_request["requested_by"] != session.get("user_id"):
                flash("You can delete only your own Office Day requests.", "error")
                return redirect(url_for("sales_dashboard"))
            if office_request["status"] != "Pending HR Approval":
                flash(
                    "Only pending Office Day requests can be deleted by Sales.",
                    "error"
                )
                return redirect(url_for("sales_dashboard"))

        connection.execute(
            "DELETE FROM office_day_requests WHERE id = ?",
            (request_id,)
        )
        connection.commit()
        flash("Office Day request deleted successfully.", "success")
    finally:
        connection.close()

    if destination == "reports":
        date_from = request.form.get("date_from", "").strip()
        date_to = request.form.get("date_to", "").strip()
        return redirect(
            url_for("reports", date_from=date_from, date_to=date_to)
        )

    return redirect(url_for(destination))


@app.route("/employees")
@roles_required("ADMIN", "HR")
def employee_master():
    search = request.args.get("search", "").strip()
    company = request.args.get("company", "").strip()
    brand = request.args.get("brand", "").strip()
    store = request.args.get("store", "").strip()
    status = request.args.get("status", "").strip()

    connection = get_db_connection()
    query = "SELECT * FROM employees WHERE 1=1"
    params = []

    if search:
        like_value = f"%{search}%"
        query += """
            AND (
                LOWER(COALESCE(company_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(last_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(first_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(middle_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(position, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(brand, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(store_assignment, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(date_hired, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(status, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(remarks, '')) LIKE LOWER(?)
            )
        """
        params.extend([like_value] * 10)

    if company:
        query += " AND LOWER(company_name)=LOWER(?)"
        params.append(company)
    if brand:
        query += " AND LOWER(brand)=LOWER(?)"
        params.append(brand)
    if store:
        query += " AND LOWER(store_assignment)=LOWER(?)"
        params.append(store)
    if status in ("Active", "Inactive"):
        query += " AND status=?"
        params.append(status)

    query += " ORDER BY last_name, first_name, middle_name"
    employees = connection.execute(query, params).fetchall()

    brands = connection.execute("""
        SELECT brand_name FROM brands WHERE status='Active'
        ORDER BY brand_name
    """).fetchall()
    stores = connection.execute("""
        SELECT store_name FROM stores WHERE status='Active'
        ORDER BY store_name
    """).fetchall()
    company_rows = connection.execute("""
        SELECT company_name FROM companies WHERE status='Active'
        ORDER BY company_name
    """).fetchall()
    connection.close()

    return render_template(
        "employee_master.html",
        employees=employees,
        brands=brands,
        stores=stores,
        companies=[row["company_name"] for row in company_rows],
        filters={"search": search, "company": company, "brand": brand, "store": store, "status": status}
    )

@app.route("/employees/save", methods=["POST"])
@roles_required("ADMIN", "HR")
def save_employee():
    employee_id = request.form.get("employee_id", "").strip()
    fields = {
        "company_name": request.form.get("company_name", "").strip(),
        "last_name": request.form.get("last_name", "").strip(),
        "first_name": request.form.get("first_name", "").strip(),
        "middle_name": request.form.get("middle_name", "").strip(),
        "position": request.form.get("position", "").strip(),
        "brand": request.form.get("brand", "").strip(),
        "store_assignment": request.form.get("store_assignment", "").strip(),
        "date_hired": request.form.get("date_hired", "").strip(),
        "status": request.form.get("status", "Active").strip(),
        "remarks": request.form.get("remarks", "").strip()
    }

    required = [
        "company_name", "last_name", "first_name",
        "position", "brand", "store_assignment"
    ]
    if any(not fields[name] for name in required):
        flash("Please complete all required employee fields.", "error")
        return redirect(url_for("employee_master"))

    if fields["status"] not in ("Active", "Inactive"):
        fields["status"] = "Active"

    connection = get_db_connection()
    if employee_id:
        connection.execute("""
            UPDATE employees SET
                company_name=?, last_name=?, first_name=?, middle_name=?,
                position=?, brand=?, store_assignment=?, date_hired=?,
                status=?, remarks=?
            WHERE id=?
        """, (
            fields["company_name"], fields["last_name"], fields["first_name"],
            fields["middle_name"], fields["position"], fields["brand"],
            fields["store_assignment"], fields["date_hired"] or None,
            fields["status"], fields["remarks"], employee_id
        ))
        message = "Employee record updated successfully."
    else:
        connection.execute("""
            INSERT INTO employees (
                company_name, last_name, first_name, middle_name,
                position, brand, store_assignment, date_hired,
                status, remarks
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            fields["company_name"], fields["last_name"], fields["first_name"],
            fields["middle_name"], fields["position"], fields["brand"],
            fields["store_assignment"], fields["date_hired"] or None,
            fields["status"], fields["remarks"]
        ))
        message = "Employee added successfully."

    connection.commit()
    connection.close()
    flash(message, "success")
    return redirect(url_for("employee_master"))


@app.route("/employees/<int:employee_id>/json")
@roles_required("ADMIN", "HR")
def employee_json(employee_id):
    connection = get_db_connection()
    row = connection.execute(
        "SELECT * FROM employees WHERE id=?", (employee_id,)
    ).fetchone()
    connection.close()
    if row is None:
        return jsonify({"error": "Employee not found"}), 404
    return jsonify(dict(row))


@app.route("/employees/<int:employee_id>/delete", methods=["POST"])
@roles_required("ADMIN", "HR")
def delete_employee(employee_id):
    connection = get_db_connection()
    try:
        employee = connection.execute(
            """
            SELECT id, last_name, first_name
            FROM employees
            WHERE id = ?
            """,
            (employee_id,)
        ).fetchone()

        if employee is None:
            flash("Employee record not found.", "error")
            return redirect(url_for("employee_master"))

        linked_records = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM office_day_requests
            WHERE employee_id = ?
            """,
            (employee_id,)
        ).fetchone()["total"]

        if linked_records:
            flash(
                "This employee has linked Office Day records and cannot be "
                "permanently deleted. Set the employee status to Inactive instead.",
                "error"
            )
            return redirect(url_for("employee_master"))

        connection.execute(
            "DELETE FROM employees WHERE id = ?",
            (employee_id,)
        )
        connection.commit()
        flash(
            f'Employee "{employee["last_name"]}, {employee["first_name"]}" '
            "deleted successfully.",
            "success"
        )
    finally:
        connection.close()

    return redirect(url_for("employee_master"))


@app.route("/employees/import", methods=["POST"])
@roles_required("ADMIN", "HR")
def import_employees():
    uploaded_file = request.files.get("employee_csv")
    if not uploaded_file or not uploaded_file.filename.lower().endswith(".csv"):
        flash("Please choose a valid employee CSV file.", "error")
        return redirect(url_for("employee_master"))

    try:
        stream = io.StringIO(
            uploaded_file.stream.read().decode("utf-8-sig"),
            newline=""
        )
        reader = csv.DictReader(stream)
    except UnicodeDecodeError:
        flash("Please save the CSV using UTF-8 format.", "error")
        return redirect(url_for("employee_master"))

    aliases = {
        "company": "company_name",
        "company name": "company_name",
        "last name": "last_name",
        "first name": "first_name",
        "middle name": "middle_name",
        "position": "position",
        "brand": "brand",
        "store": "store_assignment",
        "store assignment": "store_assignment",
        "date hired": "date_hired",
        "status": "status",
        "remarks": "remarks"
    }

    added = updated = skipped = 0
    connection = get_db_connection()
    try:
        for raw in reader:
            normalized = {}
            for key, value in raw.items():
                mapped = aliases.get((key or "").strip().lower())
                if mapped:
                    normalized[mapped] = (value or "").strip()

            required = [
                "company_name", "last_name", "first_name",
                "position", "brand", "store_assignment"
            ]
            if any(not normalized.get(field) for field in required):
                skipped += 1
                continue

            status_value = (normalized.get("status") or "Active").title()
            if status_value not in ("Active", "Inactive"):
                status_value = "Active"

            existing = connection.execute("""
                SELECT id
                FROM employees
                WHERE LOWER(last_name)=LOWER(?)
                  AND LOWER(first_name)=LOWER(?)
                ORDER BY id
                LIMIT 1
            """, (
                normalized["last_name"],
                normalized["first_name"]
            )).fetchone()

            if existing:
                connection.execute("""
                    UPDATE employees
                    SET company_name=?,
                        last_name=?,
                        first_name=?,
                        middle_name=?,
                        position=?,
                        brand=?,
                        store_assignment=?,
                        date_hired=?,
                        status=?,
                        remarks=?
                    WHERE id=?
                """, (
                    normalized["company_name"],
                    normalized["last_name"],
                    normalized["first_name"],
                    normalized.get("middle_name", ""),
                    normalized["position"],
                    normalized["brand"],
                    normalized["store_assignment"],
                    normalized.get("date_hired") or None,
                    status_value,
                    normalized.get("remarks", ""),
                    existing["id"]
                ))
                updated += 1
            else:
                connection.execute("""
                    INSERT INTO employees (
                        company_name, last_name, first_name, middle_name,
                        position, brand, store_assignment, date_hired,
                        status, remarks
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    normalized["company_name"],
                    normalized["last_name"],
                    normalized["first_name"],
                    normalized.get("middle_name", ""),
                    normalized["position"],
                    normalized["brand"],
                    normalized["store_assignment"],
                    normalized.get("date_hired") or None,
                    status_value,
                    normalized.get("remarks", "")
                ))
                added += 1

        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

    flash(
        f"Employee import completed: {added} added, {updated} updated, {skipped} skipped.",
        "success"
    )

    return redirect(url_for("employee_master"))


@app.route("/employees/export")
@roles_required("ADMIN", "HR")
def export_employees():
    search = request.args.get("search", "").strip()
    company = request.args.get("company", "").strip()
    brand = request.args.get("brand", "").strip()
    store = request.args.get("store", "").strip()
    status = request.args.get("status", "").strip()

    query = """
        SELECT company_name, last_name, first_name, middle_name,
               position, brand, store_assignment, date_hired,
               status, remarks
        FROM employees
        WHERE 1=1
    """
    params = []

    if search:
        like_value = f"%{search}%"
        query += """
            AND (
                LOWER(COALESCE(company_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(last_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(first_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(middle_name, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(position, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(brand, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(store_assignment, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(date_hired, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(status, '')) LIKE LOWER(?)
                OR LOWER(COALESCE(remarks, '')) LIKE LOWER(?)
            )
        """
        params.extend([like_value] * 10)

    if company:
        query += " AND LOWER(company_name)=LOWER(?)"
        params.append(company)
    if brand:
        query += " AND LOWER(brand)=LOWER(?)"
        params.append(brand)
    if store:
        query += " AND LOWER(store_assignment)=LOWER(?)"
        params.append(store)
    if status in ("Active", "Inactive"):
        query += " AND status=?"
        params.append(status)

    query += " ORDER BY last_name, first_name"

    connection = get_db_connection()
    rows = connection.execute(query, params).fetchall()
    connection.close()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Company Name", "Last Name", "First Name", "Middle Name",
        "Position", "Brand", "Store Assignment", "Date Hired",
        "Status", "Remarks"
    ])
    for row in rows:
        writer.writerow(list(row))

    return Response(
        "\ufeff" + output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=employee_master_filtered.csv"}
    )

@app.route("/companies")
@roles_required("ADMIN", "HR")
def company_master():
    connection = get_db_connection()
    companies = connection.execute(
        "SELECT * FROM companies ORDER BY company_name"
    ).fetchall()
    connection.close()
    return render_template("company_master.html", companies=companies)


@app.route("/companies/save", methods=["POST"])
@roles_required("ADMIN", "HR")
def save_company():
    record_id = request.form.get("record_id", "").strip()
    name = request.form.get("company_name", "").strip()
    status = request.form.get("status", "Active").strip()
    remarks = request.form.get("remarks", "").strip()

    if not name:
        flash("Company name is required.", "error")
        return redirect(url_for("company_master"))

    if status not in ("Active", "Inactive"):
        status = "Active"

    connection = get_db_connection()
    try:
        duplicate = connection.execute(
            """
            SELECT id
            FROM companies
            WHERE LOWER(company_name)=LOWER(?)
              AND id <> ?
            """,
            (name, int(record_id) if record_id else 0)
        ).fetchone()

        if duplicate:
            flash("Company name already exists.", "error")
            return redirect(url_for("company_master"))

        if record_id:
            current = connection.execute(
                "SELECT id, company_name FROM companies WHERE id=?",
                (record_id,)
            ).fetchone()

            if current is None:
                flash("Company record not found.", "error")
                return redirect(url_for("company_master"))

            if current["company_name"].strip().upper() != name:
                employee_count = connection.execute(
                    """
                    SELECT COUNT(*) AS total
                    FROM employees
                    WHERE LOWER(company_name)=LOWER(?)
                    """,
                    (current["company_name"],)
                ).fetchone()["total"]

                request_count = connection.execute(
                    """
                    SELECT COUNT(*) AS total
                    FROM office_day_requests
                    WHERE LOWER(COALESCE(company_name, ''))=LOWER(?)
                    """,
                    (current["company_name"],)
                ).fetchone()["total"]

                if employee_count or request_count:
                    flash(
                        "This company is already linked to employee or Office Day records. "
                        "To preserve history, its name cannot be changed. You may update "
                        "its Status or Remarks instead.",
                        "error"
                    )
                    return redirect(url_for("company_master"))

            connection.execute("""
                UPDATE companies
                SET company_name=?, status=?, remarks=?
                WHERE id=?
            """, (name, status, remarks, record_id))
            message = "Company updated successfully."
        else:
            connection.execute("""
                INSERT INTO companies (company_name, status, remarks)
                VALUES (?, ?, ?)
            """, (name, status, remarks))
            message = "Company added successfully."

        connection.commit()
        flash(message, "success")
    except (sqlite3.IntegrityError, PostgresIntegrityError):
        connection.rollback()
        flash("Company name already exists.", "error")
    finally:
        connection.close()

    return redirect(url_for("company_master"))


@app.route("/companies/<int:company_id>/delete", methods=["POST"])
@roles_required("ADMIN", "HR")
def delete_company(company_id):
    connection = get_db_connection()
    try:
        company = connection.execute(
            "SELECT id, company_name FROM companies WHERE id=?",
            (company_id,)
        ).fetchone()

        if company is None:
            flash("Company record not found.", "error")
            return redirect(url_for("company_master"))

        employee_count = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM employees
            WHERE LOWER(company_name)=LOWER(?)
            """,
            (company["company_name"],)
        ).fetchone()["total"]

        request_count = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM office_day_requests
            WHERE LOWER(COALESCE(company_name, ''))=LOWER(?)
            """,
            (company["company_name"],)
        ).fetchone()["total"]

        if employee_count or request_count:
            flash(
                "This company is linked to employee or Office Day records and "
                "cannot be permanently deleted. Set it to Inactive instead.",
                "error"
            )
            return redirect(url_for("company_master"))

        connection.execute("DELETE FROM companies WHERE id=?", (company_id,))
        connection.commit()
        flash(
            f'Company "{company["company_name"]}" deleted successfully.',
            "success"
        )
    finally:
        connection.close()

    return redirect(url_for("company_master"))


@app.route("/companies/export")
@roles_required("ADMIN", "HR")
def export_companies():
    return export_simple_master(
        "companies", ["company_name", "status", "remarks"],
        ["Company Name", "Status", "Remarks"], "company_master.csv"
    )


@app.route("/stores")
@roles_required("ADMIN", "HR")
def store_master():
    connection = get_db_connection()
    stores = connection.execute(
        "SELECT * FROM stores ORDER BY store_name"
    ).fetchall()
    connection.close()
    return render_template("store_master.html", stores=stores)


@app.route("/stores/save", methods=["POST"])
@roles_required("ADMIN", "HR")
def save_store():
    record_id = request.form.get("record_id", "").strip()
    name = request.form.get("store_name", "").strip().upper()
    status = request.form.get("status", "Active").strip()
    remarks = request.form.get("remarks", "").strip()

    if not name:
        flash("Store name is required.", "error")
        return redirect(url_for("store_master"))

    connection = get_db_connection()
    try:
        if record_id:
            connection.execute("""
                UPDATE stores SET store_name=?, status=?, remarks=? WHERE id=?
            """, (name, status, remarks, record_id))
            message = "Store updated successfully."
        else:
            connection.execute("""
                INSERT INTO stores (store_name, status, remarks)
                VALUES (?, ?, ?)
            """, (name, status, remarks))
            message = "Store added successfully."
        connection.commit()
        flash(message, "success")
    except (sqlite3.IntegrityError, PostgresIntegrityError):
        connection.rollback()
        flash("Store name already exists.", "error")
    finally:
        connection.close()
    return redirect(url_for("store_master"))


@app.route("/stores/<int:store_id>/delete", methods=["POST"])
@roles_required("ADMIN", "HR")
def delete_store(store_id):
    connection = get_db_connection()
    try:
        store = connection.execute(
            "SELECT id, store_name FROM stores WHERE id = ?",
            (store_id,)
        ).fetchone()

        if store is None:
            flash("Store record not found.", "error")
            return redirect(url_for("store_master"))

        employee_count = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM employees
            WHERE LOWER(store_assignment)=LOWER(?)
            """,
            (store["store_name"],)
        ).fetchone()["total"]

        request_count = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM office_day_requests
            WHERE LOWER(store_assignment)=LOWER(?)
            """,
            (store["store_name"],)
        ).fetchone()["total"]

        if employee_count or request_count:
            flash(
                "This store is linked to employee or Office Day records and "
                "cannot be permanently deleted. Set it to Inactive instead.",
                "error"
            )
            return redirect(url_for("store_master"))

        connection.execute("DELETE FROM stores WHERE id = ?", (store_id,))
        connection.commit()
        flash(f'Store "{store["store_name"]}" deleted successfully.', "success")
    finally:
        connection.close()

    return redirect(url_for("store_master"))


@app.route("/stores/import", methods=["POST"])
@roles_required("ADMIN", "HR")
def import_stores():
    uploaded_file = request.files.get("store_csv")
    if not uploaded_file or not uploaded_file.filename.lower().endswith(".csv"):
        flash("Please choose a valid store CSV file.", "error")
        return redirect(url_for("store_master"))

    try:
        stream = io.StringIO(
            uploaded_file.stream.read().decode("utf-8-sig"),
            newline=""
        )
        reader = csv.DictReader(stream)
    except UnicodeDecodeError:
        flash("Please save the CSV using UTF-8 format.", "error")
        return redirect(url_for("store_master"))

    aliases = {
        "store": "store_name",
        "store name": "store_name",
        "status": "status",
        "remarks": "remarks"
    }

    added = skipped = 0
    connection = get_db_connection()
    try:
        for raw in reader:
            normalized = {}
            for key, value in raw.items():
                mapped = aliases.get((key or "").strip().lower())
                if mapped:
                    normalized[mapped] = (value or "").strip()

            store_name = normalized.get("store_name", "").upper()
            if not store_name:
                skipped += 1
                continue

            status = normalized.get("status", "Active").title()
            if status not in ("Active", "Inactive"):
                status = "Active"

            duplicate = connection.execute(
                "SELECT id FROM stores WHERE LOWER(store_name)=LOWER(?)",
                (store_name,)
            ).fetchone()
            if duplicate:
                skipped += 1
                continue

            connection.execute(
                "INSERT INTO stores (store_name, status, remarks) VALUES (?, ?, ?)",
                (store_name, status, normalized.get("remarks", ""))
            )
            added += 1

        connection.commit()
    finally:
        connection.close()

    flash(
        f"Store import completed: {added} added, {skipped} skipped.",
        "success"
    )
    return redirect(url_for("store_master"))


@app.route("/stores/export")
@roles_required("ADMIN", "HR")
def export_stores():
    return export_simple_master(
        "stores", ["store_name", "status", "remarks"],
        ["Store Name", "Status", "Remarks"], "store_master.csv"
    )


@app.route("/brands")
@roles_required("ADMIN", "HR")
def brand_master():
    connection = get_db_connection()
    brands = connection.execute(
        "SELECT * FROM brands ORDER BY brand_name"
    ).fetchall()
    connection.close()
    return render_template("brand_master.html", brands=brands)


@app.route("/brands/save", methods=["POST"])
@roles_required("ADMIN", "HR")
def save_brand():
    record_id = request.form.get("record_id", "").strip()
    name = request.form.get("brand_name", "").strip().upper()
    status = request.form.get("status", "Active").strip()
    remarks = request.form.get("remarks", "").strip()

    if not name:
        flash("Brand name is required.", "error")
        return redirect(url_for("brand_master"))

    connection = get_db_connection()
    try:
        if record_id:
            connection.execute("""
                UPDATE brands SET brand_name=?, status=?, remarks=? WHERE id=?
            """, (name, status, remarks, record_id))
            message = "Brand updated successfully."
        else:
            connection.execute("""
                INSERT INTO brands (brand_name, status, remarks)
                VALUES (?, ?, ?)
            """, (name, status, remarks))
            message = "Brand added successfully."
        connection.commit()
        flash(message, "success")
    except (sqlite3.IntegrityError, PostgresIntegrityError):
        connection.rollback()
        flash("Brand name already exists.", "error")
    finally:
        connection.close()
    return redirect(url_for("brand_master"))


@app.route("/brands/<int:brand_id>/delete", methods=["POST"])
@roles_required("ADMIN", "HR")
def delete_brand(brand_id):
    connection = get_db_connection()
    try:
        brand = connection.execute(
            "SELECT id, brand_name FROM brands WHERE id = ?",
            (brand_id,)
        ).fetchone()

        if brand is None:
            flash("Brand record not found.", "error")
            return redirect(url_for("brand_master"))

        employee_count = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM employees
            WHERE LOWER(brand)=LOWER(?)
            """,
            (brand["brand_name"],)
        ).fetchone()["total"]

        request_count = connection.execute(
            """
            SELECT COUNT(*) AS total
            FROM office_day_requests
            WHERE LOWER(brand)=LOWER(?)
            """,
            (brand["brand_name"],)
        ).fetchone()["total"]

        if employee_count or request_count:
            flash(
                "This brand is linked to employee or Office Day records and "
                "cannot be permanently deleted. Set it to Inactive instead.",
                "error"
            )
            return redirect(url_for("brand_master"))

        connection.execute("DELETE FROM brands WHERE id = ?", (brand_id,))
        connection.commit()
        flash(f'Brand "{brand["brand_name"]}" deleted successfully.', "success")
    finally:
        connection.close()

    return redirect(url_for("brand_master"))


@app.route("/brands/export")
@roles_required("ADMIN", "HR")
def export_brands():
    return export_simple_master(
        "brands", ["brand_name", "status", "remarks"],
        ["Brand Name", "Status", "Remarks"], "brand_master.csv"
    )


def export_simple_master(table, columns, headers, filename):
    connection = get_db_connection()
    allowed = {"stores", "brands", "companies"}
    if table not in allowed:
        connection.close()
        raise ValueError("Invalid table")
    rows = connection.execute(
        f"SELECT {', '.join(columns)} FROM {table} ORDER BY {columns[0]}"
    ).fetchall()
    connection.close()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(list(row))

    return Response(
        "\ufeff" + output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.route("/users")
@roles_required("ADMIN")
def user_management():
    connection = get_db_connection()
    users = connection.execute("""
        SELECT id, full_name, username, role, status, created_at
        FROM users ORDER BY full_name
    """).fetchall()
    brands = connection.execute("""
        SELECT brand_name FROM brands
        WHERE status='Active'
        ORDER BY brand_name
    """).fetchall()
    assignment_rows = connection.execute("""
        SELECT user_id, brand_name
        FROM user_brands
        ORDER BY brand_name
    """).fetchall()
    connection.close()

    user_brand_map = {}
    for row in assignment_rows:
        user_brand_map.setdefault(str(row["user_id"]), []).append(row["brand_name"])

    return render_template(
        "user_management.html",
        users=users,
        brands=brands,
        user_brand_map=user_brand_map
    )

@app.route("/users/save", methods=["POST"])
@roles_required("ADMIN")
def save_user():
    user_id = request.form.get("user_id", "").strip()
    full_name = request.form.get("full_name", "").strip()
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "").strip().upper()
    status = request.form.get("status", "Active").strip()
    assigned_brands = [b.strip() for b in request.form.getlist("assigned_brands") if b.strip()]

    if not full_name or not username or role not in ("ADMIN", "HR", "SALES"):
        flash("Please complete all required user fields.", "error")
        return redirect(url_for("user_management"))

    if role == "SALES" and not assigned_brands:
        flash("Please assign at least one brand to a Sales user.", "error")
        return redirect(url_for("user_management"))

    if not user_id and len(password) < 8:
        flash("New account password must have at least 8 characters.", "error")
        return redirect(url_for("user_management"))

    connection = get_db_connection()
    try:
        if role == "SALES":
            active_brand_rows = connection.execute("""
                SELECT brand_name FROM brands WHERE status='Active'
            """).fetchall()
            active_brands = {row["brand_name"].strip().lower(): row["brand_name"] for row in active_brand_rows}
            normalized_brands = []
            for brand in assigned_brands:
                canonical = active_brands.get(brand.lower())
                if canonical and canonical not in normalized_brands:
                    normalized_brands.append(canonical)
            assigned_brands = normalized_brands
            if not assigned_brands:
                flash("Please assign at least one valid active brand to a Sales user.", "error")
                return redirect(url_for("user_management"))
        else:
            assigned_brands = []

        if user_id:
            if password:
                if len(password) < 8:
                    flash("Password must have at least 8 characters.", "error")
                    return redirect(url_for("user_management"))
                connection.execute("""
                    UPDATE users
                    SET full_name=?, username=?, password=?,
                        role=?, status=?
                    WHERE id=?
                """, (
                    full_name, username, generate_password_hash(password),
                    role, status, user_id
                ))
            else:
                connection.execute("""
                    UPDATE users
                    SET full_name=?, username=?, role=?, status=?
                    WHERE id=?
                """, (full_name, username, role, status, user_id))
            replace_user_brands(connection, int(user_id), assigned_brands)
            message = "User account updated successfully."
        else:
            cursor = connection.execute("""
                INSERT INTO users (
                    full_name, username, password, role, status
                )
                VALUES (?, ?, ?, ?, ?)
            """, (
                full_name, username, generate_password_hash(password),
                role, status
            ))
            if USING_POSTGRES:
                new_user = connection.execute(
                    "SELECT id FROM users WHERE LOWER(username)=LOWER(?)",
                    (username,)
                ).fetchone()
                new_user_id = new_user["id"]
            else:
                new_user_id = cursor.lastrowid
            replace_user_brands(connection, new_user_id, assigned_brands)
            message = "User account created successfully."
        connection.commit()
        flash(message, "success")
    except (sqlite3.IntegrityError, PostgresIntegrityError):
        connection.rollback()
        flash("Username already exists.", "error")
    finally:
        connection.close()
    return redirect(url_for("user_management"))


@app.route("/users/<int:user_id>/delete", methods=["POST"])
@roles_required("ADMIN")
def delete_user(user_id):
    if user_id == session.get("user_id"):
        flash("You cannot delete the account you are currently using.", "error")
        return redirect(url_for("user_management"))

    connection = get_db_connection()
    try:
        user = connection.execute(
            "SELECT id, full_name, role FROM users WHERE id = ?",
            (user_id,)
        ).fetchone()

        if user is None:
            flash("User account not found.", "error")
            return redirect(url_for("user_management"))

        if user["role"] == "ADMIN":
            admin_count = connection.execute(
                "SELECT COUNT(*) AS total FROM users WHERE role = 'ADMIN'"
            ).fetchone()["total"]
            if admin_count <= 1:
                flash("The last Admin account cannot be deleted.", "error")
                return redirect(url_for("user_management"))

        connection.execute("DELETE FROM user_brands WHERE user_id = ?", (user_id,))
        connection.execute("DELETE FROM users WHERE id = ?", (user_id,))
        connection.commit()
        flash(f'User account "{user["full_name"]}" deleted successfully.', "success")
    except (sqlite3.IntegrityError, PostgresIntegrityError):
        connection.rollback()
        flash(
            "This account has linked records and cannot be permanently deleted. "
            "Set its status to Inactive instead.",
            "error"
        )
    finally:
        connection.close()

    return redirect(url_for("user_management"))


@app.route("/reports", methods=["GET"])
@login_required
def reports():
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    rows = []

    if date_from and date_to:
        connection = get_db_connection()
        query = """
            SELECT * FROM office_day_requests
            WHERE requested_date BETWEEN ? AND ?
        """
        params = [date_from, date_to]

        if session["role"] == "SALES":
            query += " AND requested_by = ?"
            params.append(session["user_id"])

        query += " ORDER BY requested_date, employee_name"
        rows = connection.execute(query, params).fetchall()
        connection.close()

    return render_template(
        "reports.html",
        rows=rows,
        date_from=date_from,
        date_to=date_to
    )


initialize_database()

if __name__ == "__main__":
    app.run(debug=True)
